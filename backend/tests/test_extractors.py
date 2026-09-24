from pathlib import Path

import fitz
from docx import Document
from PIL import Image

from app.services import ocr_service
from app.services.docx_extractor import extract_docx_text
from app.services.file_processor import extract_text
from app.services.pdf_extractor import extract_pdf_text
from app.services.text_cleaner import clean_text


def test_pdf_extractor_reads_digital_text(tmp_path: Path) -> None:
    pdf_path = tmp_path / "paper.pdf"
    with fitz.open() as document:
        page = document.new_page()
        page.insert_text((72, 72), "A research paper title and abstract with enough text.")
        document.save(pdf_path)

    result = extract_pdf_text(pdf_path)

    assert "research paper title" in result["raw_text"]
    assert result["extraction_method"] == "pymupdf"
    assert result["needs_ocr"] is False


def test_pdf_extractor_marks_image_only_pdf_for_ocr(tmp_path: Path) -> None:
    pdf_path = tmp_path / "scanned.pdf"
    with fitz.open() as document:
        document.new_page()
        document.save(pdf_path)

    result = extract_pdf_text(pdf_path)

    assert result["raw_text"] == ""
    assert result["needs_ocr"] is True


def test_file_processor_ocr_reads_scanned_pdf(tmp_path: Path, monkeypatch) -> None:
    pdf_path = tmp_path / "scanned.pdf"
    with fitz.open() as document:
        document.new_page()
        document.new_page()
        document.save(pdf_path)
    monkeypatch.setattr(ocr_service.pytesseract, "image_to_string", lambda image: "OCR page")

    result = extract_text(pdf_path)

    assert result["raw_text"] == "OCR page\n\nOCR page"
    assert result["page_count"] == 2
    assert result["extraction_method"] == "tesseract-pdf"


def test_docx_extractor_reads_paragraphs_and_tables(tmp_path: Path) -> None:
    docx_path = tmp_path / "paper.docx"
    document = Document()
    document.add_paragraph("A DOCX research title")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Journal"
    table.cell(0, 1).text = "Example Journal"
    document.save(docx_path)

    result = extract_docx_text(docx_path)

    assert "A DOCX research title" in result["raw_text"]
    assert "Journal | Example Journal" in result["raw_text"]


def test_image_extractor_uses_ocr_service(tmp_path: Path, monkeypatch) -> None:
    image_path = tmp_path / "paper.png"
    Image.new("RGB", (100, 50), "white").save(image_path)
    monkeypatch.setattr(ocr_service.pytesseract, "image_to_string", lambda image: "OCR title")

    result = extract_text(image_path)

    assert result["raw_text"] == "OCR title"
    assert result["cleaned_text"] == "OCR title"
    assert result["extraction_method"] == "tesseract"


def test_text_cleaner_normalizes_noise_and_repeated_headers() -> None:
    raw_text = "Header  \r\n  Page 1  \r\nTitle\u00a0  with   spacing\nHeader\nHeader\n"

    cleaned = clean_text(raw_text)

    assert cleaned == "Title with spacing"