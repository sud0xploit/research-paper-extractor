from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.database.database import get_db
from app.database.models import Author, Base, Document, Publication, Verification
from app.main import app
from app.services.excel_exporter import build_batch_workbook


def test_excel_export_contains_publication_metadata() -> None:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add(
            Document(
                filename="paper.pdf",
                storage_path="uploads/paper.pdf",
                file_type="pdf",
                publication=Publication(
                    title="Exportable paper",
                    publication_type="Journal",
                    doi="10.1234/example",
                    review_status="APPROVED",
                    authors=[Author(name="Ada Lovelace", author_order=1)],
                    verifications=[Verification(verification_type="HUMAN_REVIEW", status="APPROVED")],
                ),
            )
        )
        session.commit()

    def override_get_db():
        with Session(engine) as session:
            yield session

    app.dependency_overrides[get_db] = override_get_db
    try:
        response = TestClient(app).get("/api/export/excel?status=APPROVED")
        assert response.status_code == 200
        assert response.headers["content-disposition"] == "attachment; filename=publications.xlsx"

        workbook = load_workbook(BytesIO(response.content))
        rows = list(workbook["Publications"].values)
        assert rows[0][0:3] == ("Document", "Title", "Authors")
        assert rows[1][0:3] == ("paper.pdf", "Exportable paper", "Ada Lovelace")
        assert rows[1][13:15] == ("APPROVED", "HUMAN_REVIEW: APPROVED")
        assert len(rows[0]) == len(rows[1]) == 15
    finally:
        app.dependency_overrides.pop(get_db, None)


def test_batch_workbook_keeps_extracted_text_in_its_own_column() -> None:
    workbook = load_workbook(
        build_batch_workbook(
            [
                {
                    "filename": "scan.png",
                    "status": "PROCESSED",
                    "metadata": {"title": "Detected title", "metadata_confidence": 0.8},
                    "classification": {"publication_type": "Journal", "confidence": 0.9, "status": "AUTO_APPROVED"},
                    "extraction": {"extraction_method": "tesseract", "cleaned_text": "OCR source text"},
                }
            ]
        )
    )
    rows = list(workbook["Batch Results"].values)

    assert rows[0] == (
        "File", "Title", "Authors", "Publication Type", "Confidence", "Review Status",
        "Extraction Method", "DOI", "Metadata Confidence", "Extracted Text", "Error",
    )
    assert rows[1][1] == "Detected title"
    assert rows[1][3] == "Journal"
    assert rows[1][9] == "OCR source text"
    assert rows[1][10] is None